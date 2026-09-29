"""Punto de entrada de la aplicación de escritorio."""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path

from . import APP_NAME, __version__

ASSETS = Path(__file__).resolve().parent / "assets"


def _configure_number_format() -> None:
    from PySide6.QtCore import QLocale

    from .formatting import configure

    locale = QLocale.system()
    configure(locale.decimalPoint(), locale.groupSeparator())


def create_app(argv: list[str]):
    from PySide6.QtGui import QIcon
    from PySide6.QtWidgets import QApplication

    from .gui.theme import apply_theme

    app = QApplication.instance() or QApplication(argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName(APP_NAME)
    app.setApplicationVersion(__version__)
    icon_path = ASSETS / "icon.png"
    if icon_path.exists():
        app.setWindowIcon(QIcon(str(icon_path)))
    _configure_number_format()
    apply_theme(app)
    return app


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv if argv is None else argv)
    if "--version" in argv:
        print(f"{APP_NAME} {__version__}")
        return 0
    if "--self-test" in argv:
        if sys.platform.startswith("linux") and not os.environ.get("DISPLAY"):
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        return run_self_test(argv)

    if sys.platform.startswith("win"):
        try:  # ícono propio en la barra de tareas de Windows
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("DataRF.MouserEngine")
        except Exception:  # noqa: BLE001
            pass

    app = create_app(argv)
    from PySide6.QtCore import QTimer

    from .gui.main_window import BOM_SUFFIXES, MainWindow

    window = MainWindow()
    window.show_initial()
    files =[a for a in argv[1:] if not a.startswith("-") and Path(a).suffix.lower() in BOM_SUFFIXES]
    if files:
        QTimer.singleShot(0, lambda: window.open_bom(files[0]))
    return app.exec()


# --- Autoprueba (se usa al compilar el .exe) -------------------------------------

_SELF_TEST_BOM = """Designator,Qty,Manufacturer,MPN,Description
C1 C2 C3,3,Murata,GRM188R71C104KA01D,100nF 0603
U1,1,Texas Instruments,LM358DR,Op amp
"""

_SELF_TEST_PARTS = [
    {"MouserPartNumber": "81-GRM188R71C104KA1D", "ManufacturerPartNumber": "GRM188R71C104KA01D",
     "Manufacturer": "Murata Electronics", "Description": "MLCC 0.1uF 16V X7R 0603",
     "AvailabilityInStock": "100000", "Min": "1", "Mult": "1",
     "PriceBreaks": [{"Quantity": 1, "Price": "$0.10", "Currency": "USD"},
                     {"Quantity": 10, "Price": "$0.02", "Currency": "USD"}]},
    {"MouserPartNumber": "595-LM358DR", "ManufacturerPartNumber": "LM358DR",
     "Manufacturer": "Texas Instruments", "Description": "Op Amp", "AvailabilityInStock": "0",
     "LeadTime": "6 Weeks", "Min": "1", "Mult": "1",
     "PriceBreaks": [{"Quantity": 1, "Price": "$0.42", "Currency": "USD"}]},
]


class _SelfTestResponse(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _self_test_opener(request, timeout=None):
    body = json.loads(request.data.decode("utf-8"))
    if "/cart/items/insert" in request.full_url:
        items = [{"MouserPartNumber": i["MouserPartNumber"], "Quantity": i["Quantity"], "Errors": [],
                  "InfoMessages": [], "UnitPrice": 0.1, "ExtendedPrice": round(0.1 * i["Quantity"], 2)}
                 for i in body.get("CartItems") or []]
        payload = {"Errors": [], "CartKey": "self-test-cart", "CurrencyCode": "USD", "CartItems": items}
        return _SelfTestResponse(json.dumps(payload).encode("utf-8"))
    wanted = body.get("SearchByPartRequest", {}).get("mouserPartNumber", "").upper().split("|")
    parts = [p for p in _SELF_TEST_PARTS
             if p["ManufacturerPartNumber"].upper() in wanted or p["MouserPartNumber"].upper() in wanted]
    payload = {"Errors": [], "SearchResults": {"NumberOfResult": len(parts), "Parts": parts}}
    return _SelfTestResponse(json.dumps(payload).encode("utf-8"))


def run_self_test(argv: list[str]) -> int:
    log_path = Path(tempfile.gettempdir()) / "mouser_engine_selftest.log"
    lines: list[str] = []
    try:
        app = create_app(argv)
        from PySide6.QtCore import QEventLoop, QTimer

        from openpyxl import load_workbook

        from .bom import load_table
        from .config import Settings
        from .gui.main_window import MainWindow
        from .mouser_api import MouserClient, RateLimiter

        def wait_for(condition, seconds: float = 30) -> None:
            deadline = time.monotonic() + seconds
            while not condition() and time.monotonic() < deadline:
                app.processEvents(QEventLoop.AllEvents, 100)
                time.sleep(0.02)
            app.processEvents()

        previous_config = os.environ.get("MOUSER_ENGINE_CONFIG_DIR")
        with tempfile.TemporaryDirectory() as tmp:
            # nunca tocar la configuración ni el historial reales del usuario
            os.environ["MOUSER_ENGINE_CONFIG_DIR"] = str(Path(tmp) / "config")
            bom = Path(tmp) / "bom.csv"
            bom.write_text(_SELF_TEST_BOM, encoding="utf-8")
            settings = Settings(api_key="self-test", cart_api_key="self-test", boards=10)
            factory = lambda key: MouserClient(key, opener=_self_test_opener,  # noqa: E731
                                               rate_limiter=RateLimiter(1000, 60))
            window = MainWindow(settings=settings, client_factory=factory, persist=False)
            window.interactive = False
            window.show()
            if not window.load_table(load_table(bom)):
                raise RuntimeError("no se cargó el BOM")
            wait_for(lambda: window._lookup_worker is None)
            lines.append(f"subtotal={window.summary.goods} currency={window.summary.currency}")
            if window.summary.priced != 2:
                raise RuntimeError(f"se esperaban 2 partes con precio, hay {window.summary.priced}")

            window.main_tabs.setCurrentWidget(window.scenarios)
            wait_for(lambda: bool(window.scenarios.chart.points), 15)
            window.grab()  # dibuja el gráfico de escenarios
            if window.scenarios.table.rowCount() != 7:
                raise RuntimeError("no se calcularon los escenarios de volumen")
            lines.append(f"escenarios={window.scenarios.table.rowCount()} puntos={len(window.scenarios.chart.points)}")

            out = window.export_excel(str(Path(tmp) / "cotizacion.xlsx"))
            if not out or not Path(out).exists():
                raise RuntimeError("no se generó el Excel")
            workbook = load_workbook(out)
            lines.append(f"hojas={workbook.sheetnames}")
            if "Escenarios" not in workbook.sheetnames:
                raise RuntimeError("falta la hoja Escenarios")
            workbook.close()

            entries = window.history.entries()
            if not entries:
                raise RuntimeError("no se guardó la cotización en el historial")
            restored = window.history.load(entries[0].id)
            lines.append(f"historial={len(entries)} partes={len(restored.items)}")

            if not window.create_cart():
                raise RuntimeError("no se inició la creación del carro")
            wait_for(lambda: window._cart_worker is None)
            if window.last_cart is None or window.last_cart.cart_key != "self-test-cart":
                raise RuntimeError("no se creó el carro")
            lines.append(f"carro={window.last_cart.cart_key} items={len(window.last_cart.items)}")
            window.close()
            QTimer.singleShot(0, app.quit)
            app.exec()
            if previous_config is None:
                os.environ.pop("MOUSER_ENGINE_CONFIG_DIR", None)
            else:
                os.environ["MOUSER_ENGINE_CONFIG_DIR"] = previous_config
        lines.append("SELF-TEST OK")
        code = 0
    except Exception:  # noqa: BLE001
        lines.append(traceback.format_exc())
        lines.append("SELF-TEST FAILED")
        code = 1
    text = "\n".join(lines)
    try:
        log_path.write_text(text, encoding="utf-8")
    except OSError:
        pass
    if sys.stdout is not None:
        print(text)
    return code


if __name__ == "__main__":
    raise SystemExit(main())
