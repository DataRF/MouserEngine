"""Exportación de la cotización a Excel y del carro de compra a CSV."""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from .models import (
    LEVEL_ERROR,
    LEVEL_EXCLUDED,
    LEVEL_OK,
    LEVEL_WARN,
    BomItem,
    ItemQuote,
    QuoteParams,
    QuoteSummary,
    lead_time_label,
    packaging_label,
)
from .pricing import format_breaks

_FILLS = {
    LEVEL_OK: "E3F4E8",
    LEVEL_WARN: "FFF4CC",
    LEVEL_ERROR: "FBE0E0",
    LEVEL_EXCLUDED: "EEEEEE",
}
_HEADER_FILL = "1F3A5F"
_PRICE_FMT = "#,##0.00###"
_MONEY_FMT = "#,##0.00"
_INT_FMT = "#,##0"

DETAIL_COLUMNS = [
    ("Ítem", 6), ("Líneas BOM", 10), ("Comprar", 9), ("Designadores", 22), ("MPN (BOM)", 22),
    ("Fabricante (BOM)", 18), ("Descripción (BOM)", 30), ("Cant./placa", 10), ("Cant. requerida", 12),
    ("N° Mouser", 22), ("MPN Mouser", 22), ("Fabricante Mouser", 18), ("Descripción Mouser", 36),
    ("Empaque", 14), ("Mín.", 8), ("Múlt.", 8), ("Cant. a comprar", 12), ("Precio unit.", 12),
    ("Total línea", 13), ("Moneda", 8), ("Stock Mouser", 12), ("Plazo fábrica", 13),
    ("En pedido a fábrica", 22), ("Ciclo de vida", 18), ("RoHS", 14), ("Cant. óptima", 12),
    ("Total óptimo", 13), ("Ahorro posible", 13), ("Tramos de precio", 45), ("Estado", 20),
    ("Observaciones", 70), ("Link Mouser", 16), ("Datasheet", 14), ("Consultado", 17),
]


def _num(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _detail_row(item: BomItem, q: ItemQuote) -> list:
    part = q.part
    on_order = ", ".join(f"{o.quantity} ({o.date})" if o.date else str(o.quantity)
                         for o in (part.on_order if part else []))
    return [
        item.id,
        item.rows_label,
        "Sí" if item.include else "No",
        item.designators,
        item.mpn,
        item.manufacturer,
        item.display_description,
        item.qty_per_board,
        q.required,
        part.mouser_pn if part else "",
        part.mpn if part else "",
        part.manufacturer if part else "",
        part.description if part else "",
        packaging_label(part.packaging) if part else "",
        part.min_qty if part else None,
        part.mult if part else None,
        q.buy_qty or None,
        _num(q.unit_price),
        _num(q.ext_price),
        q.currency,
        part.stock if part else None,
        lead_time_label(part.lead_time) if part else "",
        on_order,
        part.lifecycle_label if part else "",
        part.rohs if part else "",
        q.opt_qty or None,
        _num(q.opt_ext_price),
        _num(q.savings) if q.savings else None,
        format_breaks(part.price_breaks) if part else "",
        q.status,
        " ".join(q.notes),
        part.product_url if part else "",
        part.datasheet_url if part else "",
        item.looked_up_at.strftime("%d-%m-%Y %H:%M") if item.looked_up_at else "",
    ]


def export_excel(
    path: str | Path,
    items: list[BomItem],
    quotes: list[ItemQuote],
    summary: QuoteSummary,
    params: QuoteParams,
    bom_path: str = "",
    bom_grid: list[list[str]] | None = None,
    queried_at: datetime | None = None,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = Path(path)
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    thin = Side(style="thin", color="D0D7DE")
    border = Border(bottom=thin)
    currency = summary.currency or ""

    # --- Resumen ------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Cotización Mouser"
    ws["A1"].font = Font(bold=True, size=16, color=_HEADER_FILL)
    now = datetime.now()
    rows: list[tuple[str, object, str | None]] = [
        ("Generado", now.strftime("%d-%m-%Y %H:%M"), None),
        ("Precios y stock consultados", queried_at.strftime("%d-%m-%Y %H:%M") if queried_at else "—", None),
        ("Archivo BOM", Path(bom_path).name if bom_path else "—", None),
        ("Placas / equipos", params.boards, _INT_FMT),
        ("Merma general (%)", params.spares_pct, "0.0"),
        ("Merma pasivos (%)", params.passive_spares_pct, "0.0"),
        ("Optimización por tramos", "Sí" if params.optimize_breaks else "No", None),
        ("Moneda", currency or "—", None),
        (None, None, None),
        ("Partes (líneas consolidadas)", summary.items, _INT_FMT),
        ("Incluidas en la compra", summary.included, _INT_FMT),
        ("OK", summary.ok, _INT_FMT),
        ("Con advertencias", summary.warn, _INT_FMT),
        ("Con problemas", summary.error, _INT_FMT),
        ("Sin precio (no suman al total)", summary.unpriced, _INT_FMT),
        (None, None, None),
        ("Subtotal compra mínima", _num(summary.subtotal), _MONEY_FMT),
        ("Subtotal optimizado por tramos", _num(summary.optimized_subtotal), _MONEY_FMT),
        ("Ahorro posible por tramos", _num(summary.savings), _MONEY_FMT),
        ("Subtotal componentes (usado)", _num(summary.goods), _MONEY_FMT),
        ("Flete", _num(summary.freight), _MONEY_FMT),
        (f"Arancel ({params.duty_pct:g}%)", _num(summary.duty), _MONEY_FMT),
        (f"IVA ({params.vat_pct:g}%)", _num(summary.vat), _MONEY_FMT),
        ("Total estimado", _num(summary.total), _MONEY_FMT),
    ]
    if summary.total_clp is not None:
        rows.append(("Tipo de cambio (CLP)", params.fx_rate, "#,##0.00"))
        rows.append(("Total estimado en CLP", _num(summary.total_clp), "#,##0"))
    row_index = 3
    for label, value, fmt in rows:
        if label is not None:
            ws.cell(row=row_index, column=1, value=label).font = Font(bold=label.startswith("Total"))
            cell = ws.cell(row=row_index, column=2, value=value)
            if fmt:
                cell.number_format = fmt
            if label.startswith("Total"):
                cell.font = Font(bold=True)
        row_index += 1
    row_index += 1
    notes = [
        "Precios y stock según la Mouser Search API al momento de la consulta; pueden cambiar.",
        "El total de componentes no incluye envío, aranceles, IVA ni gastos de aduana salvo que "
        "se hayan ingresado en los parámetros.",
    ]
    if summary.mixed_currency:
        notes.append("ATENCIÓN: hay partes cotizadas en monedas distintas.")
    for note in notes:
        ws.cell(row=row_index, column=1, value=note).font = Font(italic=True, color="555555")
        row_index += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22

    # --- Detalle y Problemas ---------------------------------------------------
    def write_detail(sheet, pairs):
        for col, (title, width) in enumerate(DETAIL_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.row_dimensions[1].height = 30
        titles = [t for t, _ in DETAIL_COLUMNS]
        formats = {
            "Precio unit.": _PRICE_FMT, "Total línea": _MONEY_FMT, "Total óptimo": _MONEY_FMT,
            "Ahorro posible": _MONEY_FMT, "Cant./placa": _INT_FMT, "Cant. requerida": _INT_FMT,
            "Mín.": _INT_FMT, "Múlt.": _INT_FMT, "Cant. a comprar": _INT_FMT, "Stock Mouser": _INT_FMT,
            "Cant. óptima": _INT_FMT,
        }
        status_col = titles.index("Estado") + 1
        link_col = titles.index("Link Mouser") + 1
        sheet_col = titles.index("Datasheet") + 1
        for r, (item, q) in enumerate(pairs, start=2):
            values = _detail_row(item, q)
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=c, value=value)
                fmt = formats.get(titles[c - 1])
                if fmt:
                    cell.number_format = fmt
                cell.border = border
            fill = _FILLS.get(q.level)
            if fill:
                sheet.cell(row=r, column=status_col).fill = PatternFill("solid", fgColor=fill)
            for col, label in ((link_col, "Ver en Mouser"), (sheet_col, "Datasheet")):
                cell = sheet.cell(row=r, column=col)
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.value = label
                    cell.font = Font(color="0563C1", underline="single")
        sheet.freeze_panes = "B2"
        last = get_column_letter(len(DETAIL_COLUMNS))
        sheet.auto_filter.ref = f"A1:{last}{max(1, len(pairs) + 1)}"

    pairs = list(zip(items, quotes))
    write_detail(wb.create_sheet("Detalle"), pairs)
    problems = [(i, q) for i, q in pairs if q.level in (LEVEL_WARN, LEVEL_ERROR)]
    write_detail(wb.create_sheet("Problemas"), problems)

    # --- Carro Mouser ----------------------------------------------------------
    cart = wb.create_sheet("Carro Mouser")
    for col, (title, width) in enumerate(
            [("Mouser Part Number", 24), ("Quantity", 10), ("Customer Part Number", 30),
             ("Manufacturer Part Number", 26), ("Description", 50)], start=1):
        cell = cart.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cart.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(cart_rows(items, quotes), start=2):
        for c, value in enumerate(row, start=1):
            cart.cell(row=r, column=c, value=value)

    # --- BOM original ----------------------------------------------------------
    if bom_grid:
        raw = wb.create_sheet("BOM original")
        for r, row in enumerate(bom_grid, start=1):
            for c, value in enumerate(row, start=1):
                if value != "":
                    raw.cell(row=r, column=c, value=value)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def cart_rows(items: list[BomItem], quotes: list[ItemQuote]) -> list[list]:
    """Filas para cargar el carro en Mouser: código Mouser, cantidad y referencia."""
    rows = []
    for item, q in zip(items, quotes):
        if not item.include or q.level == LEVEL_EXCLUDED or q.part is None or not q.buy_qty:
            continue
        if not q.part.orderable:
            continue
        reference = (item.designators or f"Linea {item.rows_label}").replace("*", "")[:30]
        rows.append([q.part.mouser_pn, q.buy_qty, reference, q.part.mpn, q.part.description])
    return rows


def export_cart_csv(path: str | Path, items: list[BomItem], quotes: list[ItemQuote]) -> int:
    path = Path(path)
    rows = cart_rows(items, quotes)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Mouser Part Number", "Quantity", "Customer Part Number",
                         "Manufacturer Part Number", "Description"])
        writer.writerows(rows)
    return len(rows)
