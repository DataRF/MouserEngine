import csv
from datetime import datetime

from openpyxl import load_workbook

from mouser_engine.bom import build_items, load_table
from mouser_engine.export import cart_rows, export_cart_csv, export_excel
from mouser_engine.models import QuoteParams
from mouser_engine.quote import apply_lookup, queries_for, quote_all, summarize


def build(client, example_bom, **params):
    table = load_table(example_bom)
    items, _ = build_items(table)
    apply_lookup(items, client.lookup(queries_for(items)))
    p = QuoteParams(**params)
    quotes = quote_all(items, p)
    return table, items, quotes, summarize(items, quotes, p), p


def test_export_excel(client, example_bom, tmp_path):
    table, items, quotes, summary, params = build(client, example_bom, boards=10, vat_pct=19, fx_rate=900)
    path = export_excel(tmp_path / "cot.xlsx", items, quotes, summary, params, bom_path=str(example_bom),
                        bom_grid=table.grid, queried_at=datetime(2026, 9, 29, 12, 30))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Resumen", "Detalle", "Problemas", "Carro Mouser", "BOM original"]
    resumen = {row[0]: row[1] for row in wb["Resumen"].iter_rows(min_row=3, values_only=True) if row[0]}
    assert resumen["Placas / equipos"] == 10
    assert resumen["Precios y stock consultados"] == "29-09-2026 12:30"
    assert abs(resumen["Total estimado"] - float(summary.total)) < 0.001
    assert "Total estimado en CLP" in resumen
    detail = wb["Detalle"]
    headers = [c.value for c in detail[1]]
    assert headers[:3] == ["Ítem", "Líneas BOM", "Comprar"]
    rows = list(detail.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(items)
    first = dict(zip(headers, rows[0]))
    assert first["N° Mouser"] == "81-GRM188R71C104KA1D"
    assert first["Cant. a comprar"] == 80 and first["Total línea"] == 1.68
    assert first["Empaque"] == "Carrete, Cinta cortada, MouseReel"
    link_col = headers.index("Link Mouser") + 1
    assert detail.cell(row=2, column=link_col).hyperlink.target.startswith("https://www.mouser.com/")
    problems = list(wb["Problemas"].iter_rows(min_row=2, values_only=True))
    assert len(problems) == summary.warn + summary.error
    assert wb["BOM original"]["A4"].value == "Designador"


def test_cart_rows_and_csv(client, example_bom, tmp_path):
    _, items, quotes, _, _ = build(client, example_bom, boards=10)
    rows = cart_rows(items, quotes)
    pns = [r[0] for r in rows]
    assert "81-GRM188R71C104KA1D" in pns
    assert "700-MAX232CPE" not in pns          # sin precio
    assert pns.count("81-GRM188R71C104KA1D") == 1  # la línea DNP no se compra
    assert all(len(r[2]) <= 30 and "*" not in r[2] for r in rows)
    path = tmp_path / "carro.csv"
    assert export_cart_csv(path, items, quotes) == len(rows)
    with path.open(encoding="utf-8-sig") as fh:
        data = list(csv.reader(fh))
    assert data[0][:2] == ["Mouser Part Number", "Quantity"]
    assert data[1][:2] == ["81-GRM188R71C104KA1D", "80"]
